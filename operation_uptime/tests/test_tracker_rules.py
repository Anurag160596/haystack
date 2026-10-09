"""Live-round test for the inputs that only act once actuals exist (Tracker rules R1–R13, flag thresholds F1–F4).

At the base plan these inputs change nothing (no actuals are typed, no threshold is crossed), so the
one-at-a-time sweep shows them as 'no visible effect'. Here a realistic week of actuals is typed into
the Tracker; then each rule is moved through a few values and the test confirms that a Tracker decision,
re-forecast, or flag changes, with zero formula errors.

Usage: python3 test_tracker_rules.py
"""
import sys, os, json, shutil, tempfile
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import run_tests as rt

HERE = os.path.dirname(os.path.abspath(__file__))
# week-1 actuals: channel ID → (spend €, applications, qualified)
ACT = {"C09": (1449, 20, 5), "C13": (800, 12, 6), "C14": (600, 4, 0), "C17": (750, 3, 0), "C05": (500, 2, 0),
       "C15": (750, 15, 1), "C06": (400, 4, 2), "C18": (300, 3, 1), "I1a-E": (30, 12, 6), "C10": (1449, 8, 0)}   # C13 = clear performer, C10 = clear dud
# role actuals: qualified assessed, interviewed, passed, offers, signed
ROLE = {"Electricians": (30, 12, 5, 4, 2), "Mechatronics": (40, 15, 9, 7, 5), "Automation": (10, 4, 2, 1, 1), "Supervisors": (8, 3, 1, 1, 0)}
RULES = ["R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8", "R9", "R10", "R11", "R12", "R13", "F1", "F2", "F3", "F4"]
MULTS = [0.5, 2, 0.8, 1.2, 1.5, 3, 0.25, 4]
# R13 (pacing floor) only shows when a week is not already 'Full' (F4 takes precedence), so test it with slack capacity
CONTEXT = {"R13": {"G3": 80}}
SHEETS = ["Tracker", "Funnel", "Weekly", "Summary"]


def type_actuals(path):
    C = load_workbook(path, data_only=True)["Channels"]   # IDs are formulas; read their cached values
    ids = {C.cell(r, 1).value: r for r in range(5, 35)}
    wb = load_workbook(path)
    T = wb["Tracker"]
    for cid, (sp, ap, q) in ACT.items():
        r = ids[cid]  # Tracker rows mirror Channels rows
        T[f"G{r}"], T[f"H{r}"], T[f"I{r}"] = sp, ap, q
    for i, (role, v) in enumerate(ROLE.items()):
        for col, x in zip("CDEFG", v):
            T[f"{col}{40+i}"] = x
    wb.save(path)


def snap(path):
    wb = load_workbook(path, data_only=True)
    return {f"{n}!{c.coordinate}": c.value for n in SHEETS for row in wb[n].iter_rows() for c in row
            if c.value is not None and not (isinstance(c.value, str) and c.value.startswith("="))}


def main():
    tmp = tempfile.mkdtemp(prefix="uptime_trk_")
    live = os.path.join(tmp, "live.xlsx")
    shutil.copy(rt.MODEL, live); type_actuals(live)
    rc = rt.recalc(live)
    assert rc.get("total_errors", 1) == 0, rc
    base = snap(live)
    wb = load_workbook(live, data_only=True); T = wb["Tracker"]
    decisions = {T.cell(r, 1).value: T.cell(r, 24).value for r in range(5, 35) if T.cell(r, 7).value is not None}
    print("Decisions with actuals typed:", json.dumps(decisions, ensure_ascii=False))
    vals = {k: v for k, _, v, _ in __import__("test_every_input").inputs(rt.MODEL)}
    results = []
    for rule in RULES:
        hit = None
        start = live
        if rule in CONTEXT:
            start = os.path.join(tmp, f"ctx_{rule}.xlsx"); shutil.copy(live, start)
            rt.apply_edits(start, CONTEXT[rule]); rt.recalc(start)
        rbase = snap(start)
        for m in MULTS:
            p = os.path.join(tmp, f"{rule}_{m}.xlsx")
            shutil.copy(start, p)
            v = vals[rule] * m
            if rule in ("R12", "F1", "F2", "F3", "F4", "R7") and v > 1:
                v = min(v, 0.99) if rule != "F1" else v
            rt.apply_edits(p, {rule: v})
            r = rt.recalc(p)
            if r.get("total_errors", 1) != 0:
                hit = dict(value=v, ok=False, detail=f"formula errors: {r}"); break
            s = snap(p)
            diff = [k for k in set(rbase) | set(s) if rbase.get(k) != s.get(k)]
            if diff:
                ex = sorted(diff)[:3]
                hit = dict(value=v, ok=True, cells=len(diff), example=[f"{k}: {rbase.get(k)!r} → {s.get(k)!r}" for k in ex], context=CONTEXT.get(rule))
                break
        hit = hit or dict(ok=False, detail="no output changed for any tried value")
        results.append(dict(rule=rule, base=vals[rule], **hit))
        print(f"{'PASS' if hit['ok'] else 'FAIL'} {rule:4s} base {vals[rule]!s:6s} → {hit.get('value', '')!s:8.6s} "
              f"{hit.get('cells', 0):4d} cells  {(hit.get('example') or [hit.get('detail', '')])[0][:110]}", flush=True)
    json.dump(dict(decisions=decisions, results=results), open(os.path.join(HERE, "tracker_rules_results.json"), "w"), indent=1, default=str)
    p = sum(r["ok"] for r in results)
    print(f"\nRULES {p}/{len(results)} live")
    return 0 if p == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
